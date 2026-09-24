from django.test import TestCase
from django.urls import reverse
from .crypto import encrypt, lookup
from .models import Batch, Student, Submission
from django.contrib.auth import get_user_model
from io import BytesIO
from openpyxl import load_workbook


class ParentFlowTests(TestCase):
    def setUp(self):
        self.batch = Batch.objects.create(title='测试批次', fingerprint='a' * 64)
        self.student = Student.objects.create(
            batch=self.batch, source_row=3, name='测试学生', name_key=lookup('测试学生'),
            identity_cipher=encrypt({'value': '430000200001010011'}),
            identity_key=lookup('430000200001010011'), class_name='2601',
            original_phone_cipher=encrypt({'value': '13800138000'}))

    def login_parent(self):
        response = self.client.get(reverse('captcha'))
        self.assertEqual(response.status_code, 200)
        code = self.client.session['yanchi_captcha']
        response = self.client.post(reverse('login'), {'name': '测试学生', 'identity': '430000200001010011', 'captcha': code})
        self.assertRedirects(response, reverse('check'))

    def test_phone_submission_and_reopen(self):
        self.login_parent()
        response = self.client.post(reverse('check'), {'phone': '13900139000', 'result': 'issue'})
        self.assertRedirects(response, reverse('result'))
        self.student.refresh_from_db()
        self.assertEqual(self.student.status, 'submitted')
        self.assertTrue(self.student.phone_issue)
        self.assertEqual(Submission.objects.filter(student=self.student).count(), 1)
        response = self.client.post(reverse('result'))
        self.assertRedirects(response, reverse('check'))
        self.student.refresh_from_db()
        self.assertEqual(self.student.status, 'draft')

    def test_progress_filters_by_class_and_status(self):
        response = self.client.get(reverse('progress'), {'class': '2601', 'status': 'pending'})
        self.assertContains(response, '测试学生')
        self.student.status = 'submitted'; self.student.save(update_fields=['status'])
        response = self.client.get(reverse('progress'), {'class': '2601', 'status': 'pending'})
        self.assertNotContains(response, '测试学生')

    def test_staff_workspace_and_latest_phone_export(self):
        admin = get_user_model().objects.create_user(username='admin', password='secret', is_staff=True, is_superuser=True)
        Submission.objects.create(student=self.student, version=1, phone_cipher=encrypt({'value': '13900139000'}), issue=False)
        Submission.objects.create(student=self.student, version=2, phone_cipher=encrypt({'value': '13600136000'}), issue=True)
        self.student.status='submitted'; self.student.phone_issue=True; self.student.save(update_fields=['status','phone_issue'])
        self.client.force_login(admin)
        response = self.client.get(reverse('manage'), {'q': '测试学生'})
        self.assertContains(response, '测试学生')
        export = self.client.get(reverse('manage_export'), {'q': '测试学生'})
        self.assertEqual(export.status_code, 200)
        sheet = load_workbook(BytesIO(export.content), read_only=True).active
        values = [cell for row in sheet.iter_rows(values_only=True) for cell in row]
        self.assertIn('13600136000', values)
        self.assertNotIn('13900139000', values)
