import io
import base64
import tempfile
from unittest.mock import patch
from urllib.parse import unquote
import openpyxl
from PIL import Image, ImageDraw
from django.conf import settings
from django.contrib.auth import get_user_model
from django.test import TestCase, Client, override_settings
from django.db import IntegrityError, transaction
from django.urls import reverse
from .models import Batch,Student,Submission,ImportPreview
from .crypto import encrypt,decrypt,identity,digest
from .schema import FIELDS,VISIBLE,GROUPS,BY_KEY,REGION_KEYS,SENSITIVE,MACAU_REGIONS,validate_checks,valid_id
from .regions import REGION_TREE,split_region
from .services import commit_import,save_draft,submit,reopen,parse_import,export_workbook,attempt,safe_cell,public_review_issues
from .management.commands.seed_demo import demo_rows
from django.core.management import call_command
from django.core.management.base import CommandError

def test_signature():
    image=Image.new('RGB',(300,120),'white');draw=ImageDraw.Draw(image)
    draw.line([(30,80),(80,35),(125,85),(175,30),(250,75)],fill='#142f47',width=5)
    stream=io.BytesIO();image.save(stream,format='PNG')
    return 'data:image/png;base64,'+base64.b64encode(stream.getvalue()).decode()

class FlowTests(TestCase):
    TEST_SIGNATURE = test_signature()
    @classmethod
    def setUpTestData(cls):
        cls.batch=commit_import(demo_rows(),'test-batch','测试','test',demo=True)
        cls.batch.active=True;cls.batch.save()
        cls.admin=get_user_model().objects.create_superuser('tester','', 'Test-password-123456!')
    def setUp(self):
        self.student=Student.objects.get(batch=self.batch,source_row=2)
        self.client=Client();s=self.client.session;s['student_id']=str(self.student.pk);s.save()
    def all_checks(self):
        values=self.student.current()
        return {f['key']:{'result':'confirmed','value':values.get(f['key'],''),'loaded':True,'mode':'direct'} for f in VISIBLE}
    def complete(self,checks=None):
        s=save_draft(self.student.pk,self.student.revision,checks or self.all_checks(),VISIBLE)
        return submit(s.pk,s.revision,self.TEST_SIGNATURE)
    def workbook(self,rows=None):
        wb=openpyxl.load_workbook(settings.BASE_DIR/'assets/template.xlsx');ws=wb['新生1']
        for r,row in enumerate(rows or demo_rows(),2):
            for c,f in enumerate(FIELDS,1):safe_cell(ws.cell(r,c),row['values'].get(f['key'],''))
        out=io.BytesIO();wb.save(out);return out.getvalue()
    def test_schema_exact(self):
        self.assertEqual((len(FIELDS),sum(f['required'] for f in FIELDS),len(VISIBLE)),(86,25,85))
        self.assertEqual([len(g) for g in GROUPS],[15,17,27,13,13])
        self.assertEqual(sum(len(g) for g in GROUPS),len(VISIBLE))
        self.assertEqual([f['key'] for f in VISIBLE if f['readonly']],['V'])
        self.assertEqual(sum(f['required'] and not f['readonly'] for f in VISIBLE),23)
    def test_hidden_and_sensitive(self):
        response=self.client.get(reverse('step',args=[1]))
        self.assertEqual(response.status_code,200)
        self.assertNotContains(response,'学籍流水号');self.assertNotContains(response,'DEMO2026001')
        self.assertEqual(self.client.get(reverse('reveal',args=['J'])).json()['value'],'DEMO2026001')
        self.assertEqual(self.client.get(reverse('reveal',args=['A'])).status_code,403)
        self.assertIn('no-store',response.headers['Cache-Control'])

    def test_login_uses_school_identity(self):
        response=Client().get(reverse('login'))
        self.assertContains(response,'中雅实验学校')
        self.assertContains(response,'2026级新生学籍信息核对')
        self.assertContains(response,'school-logo')
        self.assertNotContains(response,self.batch.title)
        self.assertNotContains(response,'本次核对批次')
        self.assertTrue((settings.BASE_DIR/'static/school-logo.jpg').stat().st_size>10000)

    def test_option_fields_use_accessible_dropdowns(self):
        response = self.client.get(reverse('step', args=[1]))
        self.assertContains(response, 'data-long-options')
        self.assertContains(response, 'role="combobox"')
        for index, expected in [(1, 'E'), (2, 'Q'), (4, 'BO'), (5, 'CB')]:
            response = self.client.get(reverse('step', args=[index]))
            self.assertContains(response, 'data-region-picker')
            self.assertContains(response, f'province_{expected}')
            self.assertContains(response, f'city_{expected}')
            self.assertContains(response, f'district_{expected}')
            self.assertContains(response, f'name="value_{expected}"')

    def test_region_provinces_and_current_path_are_server_rendered(self):
        self.student.school_cipher=encrypt({'E':'湖南省长沙市雨花区'})
        self.student.save(update_fields=['school_cipher'])
        response=self.client.get(reverse('step',args=[1]))
        html=response.content.decode()
        picker=html.split('id="province_E"',1)[1].split('</select>',1)[0]
        self.assertEqual(picker.count('<option'),35)
        self.assertIn('>湖南省</option>',picker)
        city=html.split('id="city_E"',1)[1].split('</select>',1)[0]
        district=html.split('id="district_E"',1)[1].split('</select>',1)[0]
        self.assertIn('>长沙市</option>',city)
        self.assertIn('>雨花区</option>',district)
        self.assertNotIn('市辖区',district)

    def test_compatibility_assets_use_legacy_parseable_syntax(self):
        response=self.client.get(reverse('step',args=[1]))
        self.assertNotContains(response,'src="/xueji/static/app.js"')
        for name in ['compat-core.js','compat-regions.js','compat-sensitive.js','compat-options.js','compat-signature.js']:
            self.assertContains(response,name.replace('.js','.'))
            source=(settings.BASE_DIR/'static'/name).read_text(encoding='utf-8')
            for forbidden in ['?.','=>','replaceChildren','new Map','async function','await ']:
                self.assertNotIn(forbidden,source,name)
        signature=(settings.BASE_DIR/'static/compat-signature.js').read_text(encoding='utf-8')
        self.assertIn("'touchstart'",signature)
        self.assertIn("'pointerdown'",signature)

    def test_region_tree_covers_nationwide_template_values(self):
        values = {district['value'] for province in REGION_TREE for city in province['cities'] for district in city['districts']}
        self.assertEqual(len(REGION_TREE), 34)
        self.assertEqual(values, set(BY_KEY['Q']['options']))
        self.assertEqual(REGION_KEYS, {'E', 'F', 'Q', 'Z', 'BO', 'CB'})
        self.assertTrue(set(MACAU_REGIONS).issubset(values))
        self.assertEqual(split_region('湖南省长沙市雨花区'), ('湖南省', '长沙市', '雨花区'))
        self.assertEqual(split_region('湖北省仙桃市'), ('湖北省', '省直辖县级行政区划', '仙桃市'))
        self.assertEqual(split_region('台湾新竹县竹北市'), ('台湾省', '新竹县', '竹北市'))
        self.assertFalse(any(value.endswith('市辖区') for value in values))
        hunan = next(province for province in REGION_TREE if province['name'] == '湖南省')
        changsha = next(city for city in hunan['cities'] if city['name'] == '长沙市')
        self.assertEqual(
            [district['name'] for district in changsha['districts']],
            ['芙蓉区','天心区','岳麓区','开福区','雨花区','望城区','长沙县','浏阳市','宁乡市'],
        )

    def test_direct_edit_layout_and_confirmation_labels(self):
        response=self.client.get(reverse('step',args=[1]))
        self.assertContains(response,'name="value_B"')
        self.assertContains(response,'value="演示学生甲"')
        self.assertContains(response,'value="confirmed"')
        self.assertContains(response,'value="unconfirmed"')
        self.assertNotContains(response,'信息有误')

    def test_residency_category_guidance_is_clear_and_read_only(self):
        self.assertFalse(BY_KEY['T']['required'])
        self.assertFalse(BY_KEY['T']['readonly'])
        self.assertEqual(BY_KEY['T']['options'],[
            '常驻','蓝印','人才引进居住证','务工人员居住证',
            '在长就读的港澳台侨学生','在长就读的外国籍学生','其他',
        ])
        before=(self.student.original_cipher,self.student.school_cipher,self.student.draft_cipher,self.student.revision)
        response=self.client.get(reverse('step',args=[2]))
        self.assertContains(response,'户籍类别填写说明')
        self.assertContains(response,'aria-describedby="field-guidance-T"')
        self.assertContains(response,'不是“农业户口 / 非农业户口”的户口性质')
        self.assertContains(response,'户口所在地在长沙市时，通常选择此项')
        self.assertContains(response,'不要统一选择“其他”')
        self.student.refresh_from_db()
        self.assertEqual((self.student.original_cipher,self.student.school_cipher,self.student.draft_cipher,self.student.revision),before)
        self.assertNotContains(self.client.get(reverse('step',args=[1])),'户籍类别填写说明')

    def test_next_locates_first_validation_error(self):
        values=self.student.current();data={'revision':self.student.revision,'action':'next'}
        for field in GROUPS[0]:
            key=field['key'];data['value_'+key]=values.get(key,'');data['result_'+key]='confirmed'
            if key in {'J','BT'}:data['loaded_'+key]='1'
        data['result_B']='unconfirmed'
        response=self.client.post(reverse('step',args=[1]),data)
        self.assertEqual(response.status_code,200)
        self.assertEqual(response.context['error_count'],1)
        self.assertContains(response,'还有 1 项需要处理')
        self.assertContains(response,'data-field="B" data-first-error tabindex="-1"')
        self.assertContains(response,'请处理此项')
        self.assertContains(response,'保存并进入下一步')
    def test_anonymous_cannot_reveal_or_export(self):
        client=Client()
        for url in [reverse('step',args=[1]),reverse('reveal',args=['J']),reverse('dashboard'),reverse('export',args=[self.batch.pk,'final'])]:
            self.assertEqual(client.get(url).status_code,302)
    def test_draft_resumes_and_encrypts(self):
        checks=self.all_checks();checks['BT']={'result':'confirmed','value':'PRIVATE-GUARDIAN','loaded':True}
        s=save_draft(self.student.pk,0,checks,VISIBLE)
        self.assertEqual(s.draft['BT']['value'],'PRIVATE-GUARDIAN')
        self.assertNotIn('PRIVATE-GUARDIAN',s.draft_cipher)
        self.assertNotIn('DEMO2026001',s.original_cipher)
        self.assertNotContains(self.client.get(reverse('step',args=[4])),'PRIVATE-GUARDIAN')
        self.assertEqual(self.client.get(reverse('reveal',args=['BT'])).json()['value'],'PRIVATE-GUARDIAN')

    def test_sensitive_value_is_preserved_until_loaded(self):
        checks={'J':{'result':'confirmed','value':'','loaded':False}}
        s=save_draft(self.student.pk,0,checks,[BY_KEY['J']])
        self.assertEqual(s.draft['J']['value'],'DEMO2026001')

    def test_blank_sensitive_value_is_directly_editable(self):
        values=self.student.original;values['BT']=''
        self.student.original_cipher=encrypt(values);self.student.save(update_fields=['original_cipher'])
        response=self.client.get(reverse('step',args=[4]))
        html=response.content.decode()
        card=html.split('id="field-BT"',1)[1].split('</section>',1)[0]
        self.assertIn('name="loaded_BT" value="1"',card)
        self.assertIn('name="value_BT" value=""',card)
        self.assertNotIn('data-reveal=',card)

    def test_sensitive_reveal_form_fallback_preserves_original_and_opens_editor(self):
        original=self.student.original_cipher
        values=self.student.current();data={'revision':self.student.revision,'action':'reveal-sensitive:J'}
        for field in GROUPS[0]:
            key=field['key'];data['result_'+key]='confirmed'
            if not field['readonly']:
                data['value_'+key]=values.get(key,'')
            if key in SENSITIVE:data['loaded_'+key]='0'
        response=self.client.post(reverse('step',args=[1]),data)
        self.assertEqual(response.status_code,302)
        self.assertIn('?reveal=J#field-J',response.url)
        self.student.refresh_from_db()
        self.assertEqual(self.student.original_cipher,original)
        revealed=self.client.get(response.url)
        self.assertContains(revealed,'name="loaded_J" value="1"')
        self.assertContains(revealed,'name="value_J" value="DEMO2026001"')
        self.assertIn('no-store',revealed.headers['Cache-Control'])

    def test_region_compatibility_flow_uses_existing_encrypted_draft(self):
        original=self.student.original_cipher
        values=self.student.current();data={'revision':self.student.revision,'action':'compat-region:Q'}
        for field in GROUPS[1]:
            key=field['key'];data['result_'+key]='confirmed'
            if not field['readonly']:data['value_'+key]=values.get(key,'')
        previous_q=values['Q'];data['value_Q']=''
        response=self.client.post(reverse('step',args=[2]),data)
        self.assertRedirects(response,reverse('region_compat',args=['Q']))
        self.student.refresh_from_db()
        self.assertEqual(self.student.draft['Q']['value'],previous_q)
        page=self.client.get(reverse('region_compat',args=['Q']),{'q':'雨花区'})
        self.assertContains(page,'湖南省长沙市雨花区')
        response=self.client.post(reverse('region_compat',args=['Q']),{
            'revision':self.student.revision,'value':'湖南省长沙市雨花区',
        })
        self.assertEqual(response.status_code,302)
        self.student.refresh_from_db()
        self.assertEqual(self.student.original_cipher,original)
        self.assertEqual(self.student.draft['Q']['value'],'湖南省长沙市雨花区')
        self.assertEqual(self.student.draft['Q']['result'],'unconfirmed')
        self.assertNotIn('湖南省长沙市雨花区',self.student.draft_cipher)

    def test_region_compatibility_rejects_unknown_field_and_value(self):
        self.assertEqual(self.client.get(reverse('region_compat',args=['J'])).status_code,403)
        response=self.client.post(reverse('region_compat',args=['Q']),{
            'revision':self.student.revision,'value':'湖南省长沙市市辖区',
        })
        self.assertEqual(response.status_code,200)
        self.assertContains(response,'请选择列表中的完整行政区划')
        self.student.refresh_from_db();self.assertNotIn('Q',self.student.draft)
        revision=self.student.revision
        response=self.client.post(reverse('step',args=[2]),{'revision':revision,'action':'compat-region:J'})
        self.assertEqual(response.status_code,200)
        self.assertContains(response,'无法识别本次操作')
        self.student.refresh_from_db();self.assertEqual(self.student.revision,revision)

    def test_legacy_draft_statuses_are_compatible(self):
        self.student.draft_cipher=encrypt({'B':{'result':'incorrect','value':'旧草稿姓名'},
                                           'C':{'result':'correct','value':''},
                                           'V':{'result':'incorrect','note':'旧草稿班级问题'},
                                           'Y':{'result':'incorrect','note':'旧草稿学籍号问题','mode':'direct'}})
        self.student.status='draft';self.student.save()
        first=self.client.get(reverse('step',args=[1]))
        self.assertContains(first,'value="旧草稿姓名"')
        self.assertContains(first,'name="result_B" value="confirmed" checked')
        second=self.client.get(reverse('step',args=[2]))
        self.assertContains(second,'name="result_V" value="confirmed"')
        self.assertNotContains(second,'name="result_V" value="unconfirmed"')
        self.assertContains(second,'班级由学校统一维护，仅供查看，无需家长确认')
        self.assertContains(second,'name="value_Y" value="DEMO-NATIONAL-0001"')
        self.assertContains(second,'name="result_Y" value="unconfirmed" checked')
    def test_stale_readonly_national_id_form_preserves_value_during_deploy(self):
        data={'revision':self.student.revision,'action':'save'}
        for field in GROUPS[1]:
            data['result_'+field['key']]='confirmed'
            if field['key']!='Y' and not field['readonly']:
                data['value_'+field['key']]=self.student.current().get(field['key'],'')
        response=self.client.post(reverse('step',args=[2]),data)
        self.assertEqual(response.status_code,200)
        self.student.refresh_from_db()
        self.assertEqual(self.student.draft['Y']['value'],'DEMO-NATIONAL-0001')
        self.assertEqual(self.student.draft['Y']['result'],'confirmed')
    def test_readonly_injection_ignored(self):
        checks=self.all_checks();checks['V']={'result':'unconfirmed','value':'恶意班级','note':'请检查班级'}
        s=self.complete(checks)
        self.assertEqual(s.current()['V'],'2601');self.assertFalse(s.has_issue)
        self.assertNotIn('value',s.draft['V'])
        self.assertEqual(s.draft['V']['result'],'confirmed')
        self.assertEqual(s.draft['V']['note'],'')
        self.assertNotContains(self.client.get(reverse('result')),'学校核实')
        self.client.force_login(self.admin)
        self.assertNotContains(self.client.get(reverse('student_detail',args=[s.pk])),'请检查班级')
    def test_reopened_class_issue_is_ignored(self):
        checks=self.all_checks();checks['V']={'result':'unconfirmed','note':'请核实班级'}
        s=self.complete(checks);reopen(s.pk,self.admin);s.refresh_from_db()
        self.assertEqual(s.status,'draft');self.assertFalse(s.has_issue)
        self.client.force_login(self.admin)
        dashboard=self.client.get(reverse('dashboard'))
        self.assertEqual(dashboard.context['issues'],0)
    def test_submit_idempotent_and_edit_locked(self):
        s=self.complete();submit(s.pk,s.revision,self.TEST_SIGNATURE)
        self.assertEqual(Submission.objects.filter(student=s).count(),1)
        with self.assertRaises(ValueError):save_draft(s.pk,s.revision,{},VISIBLE)
    def test_stale_revision(self):
        save_draft(self.student.pk,0,{},GROUPS[0])
        with self.assertRaises(ValueError):save_draft(self.student.pk,0,{},GROUPS[0])
    def test_missing_blocks_submit(self):
        s=save_draft(self.student.pk,0,{},VISIBLE)
        with self.assertRaises(ValueError):submit(s.pk,s.revision,self.TEST_SIGNATURE)

    def test_signature_required_and_saved(self):
        s=save_draft(self.student.pk,0,self.all_checks(),VISIBLE)
        with self.assertRaisesMessage(ValueError,'请在签名框中完成手写签字'):
            submit(s.pk,s.revision)
        blank=Image.new('RGB',(300,120),'white');stream=io.BytesIO();blank.save(stream,format='PNG')
        blank_signature='data:image/png;base64,'+base64.b64encode(stream.getvalue()).decode()
        with self.assertRaisesMessage(ValueError,'签名笔迹过少'):
            submit(s.pk,s.revision,blank_signature)
        s=submit(s.pk,s.revision,self.TEST_SIGNATURE)
        self.assertEqual(s.submissions.get().payload['signature'],self.TEST_SIGNATURE)

    def test_confirm_and_result_show_signature_ui(self):
        response=self.client.get(reverse('confirm'))
        self.assertContains(response,'signature-pad');self.assertContains(response,'signature-input');self.assertContains(response,'手写签字')
        self.assertContains(response,'data-signature-unavailable')
        self.assertEqual(response.context['required_total'],23)
        self.complete();response=self.client.get(reverse('result'))
        self.assertContains(response,'已完成手机手写签字');self.assertContains(response,'本次电子签名');self.assertContains(response,'重新核对并修改')
    def test_name_and_id_changes_keep_original_login(self):
        checks=self.all_checks();checks['B']={'result':'confirmed','value':'更正姓名'};checks['J']={'result':'confirmed','value':'NEW-DEMO-ID','loaded':True}
        s=self.complete(checks)
        self.assertEqual(s.current()['B'],'更正姓名')
        self.assertEqual(s.identity_hash,identity('演示学生甲','DEMO2026001'))
        self.assertNotEqual(s.identity_hash,identity('更正姓名','NEW-DEMO-ID'))
    def test_reopen_retains_versions(self):
        s=self.complete();reopen(s.pk,self.admin);s.refresh_from_db()
        self.assertEqual(s.status,'draft');self.assertEqual(s.draft,{})
        s=save_draft(s.pk,s.revision,self.all_checks(),VISIBLE);submit(s.pk,s.revision,self.TEST_SIGNATURE)
        self.assertEqual(s.submissions.count(),2)
    def test_parent_can_reopen_own_submission_without_losing_history(self):
        s=self.complete();before=s.current();old_revision=s.revision
        response=self.client.post(reverse('result'),{'action':'reopen'})
        self.assertRedirects(response,reverse('step',args=[1]))
        s.refresh_from_db()
        self.assertEqual(s.status,'draft');self.assertEqual(s.draft,{})
        self.assertEqual(s.revision,old_revision+1);self.assertEqual(s.submissions.count(),1)
        self.assertEqual(s.current(),before)
    def test_blank_national_student_id_must_be_filled_and_confirmed(self):
        self.student=Student.objects.get(source_row=4,batch=self.batch)
        self.assertFalse(self.student.has_issue)
        checks=self.all_checks();checks['Y']={'result':'unconfirmed','value':'','mode':'direct'}
        self.assertIn('Y',validate_checks(self.student.current(),checks)[1])
        checks['Y']={'result':'confirmed','value':'G430100201401010001','mode':'direct'}
        s=self.complete(checks)
        self.assertEqual(s.current()['Y'],'G430100201401010001');self.assertFalse(s.has_issue)
    def test_national_student_id_parent_change_overrides_school_baseline_and_exports(self):
        school={'Y':'SCHOOL-NATIONAL-OLD','Q':'湖南省长沙市雨花区'}
        self.student.school_cipher=encrypt(school);self.student.save(update_fields=['school_cipher'])
        checks=self.all_checks();checks['Y']={'result':'confirmed','value':'G430100201401010002','mode':'direct'}
        s=self.complete(checks)
        self.assertEqual(s.school,school);self.assertEqual(s.current()['Y'],'G430100201401010002')
        self.assertEqual(s.submissions.get().payload['changes']['Y'],{'old':'SCHOOL-NATIONAL-OLD','new':'G430100201401010002'})
        workbook=openpyxl.load_workbook(io.BytesIO(export_workbook(self.batch,'final')))
        self.assertEqual(workbook['新生1']['Y2'].value,'G430100201401010002')
        self.assertEqual(workbook['新生1']['Y2'].data_type,'s')
    def test_export_preserves_template_nonrequired_and_strings(self):
        checks=self.all_checks();checks['AA']={'result':'confirmed','value':'=1+1'};s=self.complete(checks)
        wb=openpyxl.load_workbook(io.BytesIO(export_workbook(self.batch,'final')))
        self.assertEqual(wb.sheetnames,['新生1','area','nationality'])
        self.assertEqual(wb['新生1']['AA2'].value,'=1+1');self.assertEqual(wb['新生1']['AA2'].data_type,'s')
        self.assertEqual(wb['新生1']['N2'].value,'DEMO-OPTIONAL-PRESERVED')
        self.assertEqual(wb['新生1']['J2'].data_type,'s');self.assertEqual(wb['area'].sheet_state,'hidden')
        area_values={cell[0].value for cell in wb['area'].iter_rows() if cell[0].value}
        self.assertTrue(set(MACAU_REGIONS).issubset(area_values));self.assertNotIn('\u5cea泉镇',area_values)
        self.assertFalse(any(str(value).endswith('市辖区') for value in area_values))
        self.assertEqual(len(wb['新生1'].data_validations.dataValidation),49)
        self.assertEqual(wb['新生1'].max_column,86)
    def test_final_export_uses_prefetched_submissions(self):
        with self.assertNumQueries(2):
            data=export_workbook(self.batch,'final')
        self.assertTrue(data.startswith(b'PK'))
    def test_pending_export_and_changes(self):
        self.complete()
        wb=openpyxl.load_workbook(io.BytesIO(export_workbook(self.batch,'pending')))
        self.assertEqual(wb['新生1'].max_row,5)
        self.assertTrue(export_workbook(self.batch,'changes').startswith(b'PK'))
        review=openpyxl.load_workbook(io.BytesIO(export_workbook(self.batch,'review')))
        self.assertEqual(review['新生1'].max_row,6)
        self.assertEqual(review['新生1'].max_column,86)
        self.assertEqual(review['新生1']['J2'].data_type,'s')

    def test_pending_students_can_be_viewed_and_downloaded_by_class(self):
        self.complete();self.client.force_login(self.admin)
        response=self.client.get(reverse('pending_classes'),{'batch':self.batch.pk})
        self.assertEqual(response.status_code,200)
        rows={row['class_name']:row for row in response.context['class_rows']}
        self.assertEqual((rows['2601']['total'],rows['2601']['submitted'],rows['2601']['pending']),(3,1,2))
        self.assertEqual((rows['2602']['total'],rows['2602']['submitted'],rows['2602']['pending']),(2,0,2))
        self.assertContains(response,'按班查看学生核对情况')
        response=self.client.get(reverse('pending_classes'),{'batch':self.batch.pk,'class':'2601'})
        self.assertEqual(list(response.context['students'].values_list('source_row',flat=True)),[3,4])
        self.assertContains(response,'2601学生核对情况')
        self.assertContains(response,'未核对（2）')
        response=self.client.get(reverse('pending_classes'),{'class':'2601','status':'submitted'})
        self.assertEqual(list(response.context['students'].values_list('source_row',flat=True)),[2])
        response=self.client.get(reverse('pending_classes'),{'class':'2601','status':'all'})
        self.assertEqual(list(response.context['students'].values_list('source_row',flat=True)),[2,3,4])
        response=self.client.get(reverse('export',args=[self.batch.pk,'pending']),{'class':'2601'})
        self.assertEqual(response.status_code,200)
        self.assertIn('2601未完成核对学生',unquote(response.headers['Content-Disposition']))
        workbook=openpyxl.load_workbook(io.BytesIO(response.content))
        self.assertEqual(workbook['新生1'].max_row,3)

    def test_class_pending_page_is_public_without_sensitive_links(self):
        client=Client()
        response=client.get(reverse('pending_classes'),{'batch':999,'class':'2601'},HTTP_X_REAL_IP='203.0.113.10')
        self.assertEqual(response.status_code,200)
        self.assertEqual(response.context['batch'],self.batch)
        self.assertTrue(response.context['standalone'])
        self.assertContains(response,'2601学生核对情况')
        self.assertContains(response,'返回班级列表')
        self.assertNotContains(response,'各班完成情况')
        self.assertNotContains(response,'2602')
        self.assertNotContains(response,'下载本班名单')
        self.assertNotContains(response,'下载全部未完成名单')
        self.assertNotContains(response,'查看学生详情')
        self.assertNotContains(response,'学生身份证件号码')

    def test_class_pending_page_never_links_to_admin(self):
        self.client.force_login(self.admin)
        response=self.client.get(reverse('pending_classes'),{'batch':999,'class':'2601'})
        self.assertEqual(response.context['batch'],self.batch)
        for text in ['核对工作台','账户与记录','返回核对工作台','下载本班名单','查看学生详情']:
            self.assertNotContains(response,text)
        self.assertNotContains(response,'href="/xueji/manage/')

    def test_head_campus_has_separate_progress_population(self):
        head=Student.objects.get(batch=self.batch,source_row=2)
        head.progress_group='benbu';head.save(update_fields=['progress_group'])
        regular=self.client.get(reverse('pending_classes'))
        regular_2601=next(row for row in regular.context['class_rows'] if row['class_name']=='2601')
        self.assertEqual(regular_2601['total'],2)
        head_page=Client().get(reverse('pending_head_campus'),{'class':'2601'})
        self.assertEqual(head_page.context['group_label'],'本部')
        self.assertEqual(head_page.context['pending_total'],1)
        self.assertEqual(list(head_page.context['students'].values_list('pk',flat=True)),[head.pk])
        self.assertNotContains(head_page,Student.objects.get(batch=self.batch,source_row=3).name)
        self.assertContains(head_page,'本部学生总表')
        self.assertNotContains(head_page,'各班完成情况')
        self.assertNotContains(head_page,'查看本班学生')
        self.assertContains(head_page,'未核对（1）')
        self.assertContains(head_page,'已核对（0）')
        self.assertContains(head_page,'核对有误（0）')
        self.assertContains(head_page,'学籍号有误（0）')
        self.assertEqual(head_page.context['status_filter'],'all')
        self.assertEqual(head_page.context['selected_count'],1)
        pending=Client().get(reverse('pending_head_campus'),{'status':'pending'})
        self.assertEqual(list(pending.context['students'].values_list('pk',flat=True)),[head.pk])
        submitted=Client().get(reverse('pending_head_campus'),{'status':'submitted'})
        self.assertFalse(submitted.context['students'].exists())

    def test_public_progress_rosters_show_specific_issue_students(self):
        national=self.student
        national.status='draft';national.has_issue=True
        national.draft_cipher=encrypt({'Y':{'result':'unconfirmed','value':national.current()['Y'],'mode':'direct'}})
        national.save(update_fields=['status','has_issue','draft_cipher'])
        Submission.objects.create(student=national,version=1,payload_cipher=encrypt({
            'values':{},'checks':{'Y':{'result':'unconfirmed','note':''}},'changes':{},'signature':'existing-signature'}))
        self.assertEqual(public_review_issues(national),['全国学籍号待重新核对'])
        unrelated=Student.objects.get(batch=self.batch,source_row=4)
        Submission.objects.create(student=unrelated,version=1,payload_cipher=encrypt({
            'values':{},'checks':{},'changes':{},'signature':'existing-signature'}))
        self.assertEqual(public_review_issues(unrelated),[])

        regular=Client().get(reverse('pending_classes'),{'class':'2601','status':'issue'})
        self.assertContains(regular,national.name)
        self.assertContains(regular,'全国学籍号待重新核对')
        self.assertContains(regular,'核对有误（1）')
        self.assertContains(regular,'学籍号有误（1）')
        self.assertNotContains(regular,Student.objects.get(batch=self.batch,source_row=3).name)
        national_only=Client().get(reverse('pending_classes'),{'class':'2601','status':'national_issue'})
        self.assertContains(national_only,national.name)
        self.assertContains(national_only,'全国学籍号待重新核对')

        class_issue=Student.objects.get(batch=self.batch,source_row=3)
        class_issue.progress_group='benbu';class_issue.status='submitted';class_issue.has_issue=True
        class_issue.save(update_fields=['progress_group','status','has_issue'])
        Submission.objects.create(student=class_issue,version=1,payload_cipher=encrypt({
            'values':{},'checks':{'V':{'result':'unconfirmed','note':'班级有误'}},'changes':{},'signature':'existing-signature'}))
        head=Client().get(reverse('pending_head_campus'),{'status':'issue'})
        self.assertNotContains(head,class_issue.name)
        self.assertNotContains(head,'班级信息待核实')
        self.assertContains(head,'核对有误（0）')
        self.assertContains(head,'学籍号有误（0）')
        self.assertNotContains(head,national.name)
        self.assertNotContains(head,'href="/xueji/manage/')
        head_national=Client().get(reverse('pending_head_campus'),{'status':'national_issue'})
        self.assertNotContains(head_national,class_issue.name)
        self.assertContains(head_national,'本部名单没有学籍号有误学生')

    def test_assign_progress_group_command_is_atomic(self):
        students=list(Student.objects.filter(batch=self.batch).order_by('source_row')[:2])
        with tempfile.TemporaryDirectory() as directory:
            path=__import__('pathlib').Path(directory)/'本部.xlsx'
            workbook=openpyxl.Workbook();sheet=workbook.active
            sheet.append(['序号','班级编号','姓名'])
            for index,student in enumerate(students,1):sheet.append([index,student.class_name,student.name])
            workbook.save(path)
            call_command('assign_progress_group',str(path),'benbu')
            self.assertFalse(Student.objects.filter(batch=self.batch,progress_group='benbu').exists())
            call_command('assign_progress_group',str(path),'benbu','--apply')
            self.assertEqual(Student.objects.filter(batch=self.batch,progress_group='benbu').count(),2)
            sheet.append([3,'9999','不存在学生']);workbook.save(path)
            with self.assertRaises(CommandError):call_command('assign_progress_group',str(path),'benbu','--apply')
            self.assertEqual(Student.objects.filter(batch=self.batch,progress_group='benbu').count(),2)

    def test_confirmation_export_contains_every_field_and_signature(self):
        self.complete()
        wb=openpyxl.load_workbook(io.BytesIO(export_workbook(self.batch,'confirmations')))
        self.assertEqual(wb.sheetnames,['逐项确认','家长签名'])
        confirmations=wb['逐项确认'];signatures=wb['家长签名']
        self.assertEqual(confirmations.max_row,1+len(VISIBLE))
        self.assertEqual(confirmations['H2'].value,'已确认')
        self.assertEqual(confirmations['K2'].value,'否')
        self.assertEqual(signatures.max_row,2)
        self.assertEqual(len(signatures._images),1)

    def test_legacy_readonly_national_id_remains_under_school_in_history_exports(self):
        payload={'values':{},'checks':{'V':{'result':'correct'},'Y':{'result':'incorrect','note':'旧版交由学校核实'}},'changes':{},'signature':''}
        Submission.objects.create(student=self.student,version=1,payload_cipher=encrypt(payload))
        self.student.status='submitted';self.student.has_issue=True;self.student.save(update_fields=['status','has_issue'])
        changes=openpyxl.load_workbook(io.BytesIO(export_workbook(self.batch,'changes')))['修改明细']
        national_row=next(row for row in changes.iter_rows(min_row=2,values_only=True) if row[4]=='全国学籍号')
        self.assertEqual(national_row[7],'学校待处理');self.assertEqual(national_row[8],'旧版交由学校核实')
        confirmations=openpyxl.load_workbook(io.BytesIO(export_workbook(self.batch,'confirmations')))['逐项确认']
        national_row=next(row for row in confirmations.iter_rows(min_row=2,values_only=True) if row[4]=='全国学籍号')
        self.assertEqual(national_row[6],'学校');self.assertEqual(national_row[7],'未确认');self.assertEqual(national_row[10],'否')
        self.assertContains(self.client.get(reverse('result')),'重新核对并修改')
        self.assertNotContains(self.client.get(reverse('result')),'学校核实')
        self.client.force_login(self.admin)
        dashboard=self.client.get(reverse('dashboard'))
        self.assertEqual(dashboard.context['issues'],0)

    def test_staff_can_download_school_review_workbook(self):
        self.client.force_login(self.admin)
        response=self.client.get(reverse('export',args=[self.batch.pk,'review']))
        self.assertEqual(response.status_code,200)
        self.assertIn("学校自行核对表",unquote(response.headers['Content-Disposition']))
        workbook=openpyxl.load_workbook(io.BytesIO(response.content))
        self.assertEqual(workbook['新生1'].max_column,86)
        response=self.client.get(reverse('export',args=[self.batch.pk,'confirmations']))
        self.assertEqual(response.status_code,200)
        self.assertIn('家长逐项确认与签名',unquote(response.headers['Content-Disposition']))
    def test_import_roundtrip_and_duplicates(self):
        rows,fingerprint=parse_import(self.workbook());self.assertEqual(len(rows),5)
        commit_import(rows,fingerprint,'导入测试',self.admin)
        with self.assertRaises(IntegrityError),transaction.atomic():commit_import(rows,fingerprint,'重复',self.admin)
        with self.assertRaisesMessage(ValueError,'重复'):parse_import(self.workbook([demo_rows()[0],demo_rows()[0]]))
    def test_import_rejects_missing_identity_and_numeric_id(self):
        rows=demo_rows();rows[0]['values']['A']=''
        with self.assertRaisesMessage(ValueError,'不能为空'):parse_import(self.workbook(rows))
        wb=openpyxl.load_workbook(io.BytesIO(self.workbook()));wb['新生1']['J2']=123456789012345678
        out=io.BytesIO();wb.save(out)
        with self.assertRaisesMessage(ValueError,'文本格式'):parse_import(out.getvalue())
    def test_captcha_replay_and_login(self):
        c=Client();session=c.session;session['captcha']={'hash':digest('AB234'),'time':__import__('time').time()};session.save()
        r=c.post(reverse('login'),{'name':'演示学生甲','number':'DEMO2026001','captcha':'ab234'})
        self.assertEqual(r.status_code,302);self.assertEqual(c.session['student_id'],str(self.student.pk))
        c=Client();r=c.post(reverse('login'),{'name':'演示学生甲','number':'DEMO2026001','captcha':'ab234'})
        self.assertNotIn('student_id',c.session)
    def test_csrf_enforced(self):
        c=Client(enforce_csrf_checks=True)
        self.assertEqual(c.post(reverse('login'),{}).status_code,403)
    def test_demo_import_blocked(self):
        self.client.force_login(self.admin)
        self.assertContains(self.client.post(reverse('import'),{'title':'real','file':io.BytesIO(self.workbook())}),'限制真实数据')
        self.assertEqual(ImportPreview.objects.count(),0)
    def test_throttling(self):
        self.assertTrue(attempt('test',2));self.assertTrue(attempt('test',2));self.assertFalse(attempt('test',2))
    def test_closed_batch_blocks_parent(self):
        self.batch.is_open=False;self.batch.save()
        self.assertEqual(self.client.get(reverse('step',args=[1])).status_code,403)
    def test_admin_pages_and_school_update(self):
        self.student.school_cipher=encrypt({'Y':'SCHOOL-NATIONAL','Q':'湖南省长沙市雨花区'})
        self.student.save(update_fields=['school_cipher'])
        self.client.force_login(self.admin)
        for url in [reverse('dashboard'),reverse('pending_classes'),reverse('student_detail',args=[self.student.pk]),'/xueji/admin/']:
            self.assertEqual(self.client.get(url).status_code,200)
        detail=self.client.get(reverse('student_detail',args=[self.student.pk]))
        self.assertNotContains(detail,'name="Y"')
        r=self.client.post(reverse('student_detail',args=[self.student.pk]),{'action':'school','V':'2609','Y':'FORGED-NATIONAL'})
        self.assertEqual(r.status_code,302);self.student.refresh_from_db()
        self.assertEqual(self.student.current()['V'],'2609')
        self.assertEqual(self.student.school,{'Y':'SCHOOL-NATIONAL','Q':'湖南省长沙市雨花区','V':'2609'})
    def test_qr_and_template(self):
        self.client.force_login(self.admin)
        self.assertEqual(self.client.get(reverse('qr')).headers['Content-Type'],'image/png')
        self.assertEqual(self.client.get(reverse('template')).content,(settings.BASE_DIR/'assets/template.xlsx').read_bytes())
    def test_options_and_identity_validation(self):
        values=self.student.current();checks=self.all_checks();checks['G']={'result':'confirmed','value':'不在模板的民族','mode':'direct'}
        _,errors=validate_checks(values,checks);self.assertIn('G',errors)
        self.assertFalse(valid_id('123456789012345678'))

    def test_nonrequired_field_may_be_blank_and_all_template_fields_are_checked(self):
        values=self.student.current();checks=self.all_checks()
        checks['E']={'result':'confirmed','value':'','mode':'direct'}
        _,errors=validate_checks(values,checks);self.assertNotIn('E',errors)
        self.assertIn('G',validate_checks(values,{**checks,'G':{'result':'confirmed','value':'','mode':'direct'}})[1])

    def test_confirmation_rules_for_required_and_changed_optional_fields(self):
        values=self.student.current();checks=self.all_checks()
        checks['C']['result']='unconfirmed'
        self.assertIn('C',validate_checks(values,checks)[1])
        checks=self.all_checks();checks['N']['result']='unconfirmed'
        self.assertNotIn('N',validate_checks(values,checks)[1])
        checks['N']['value']='修改后的非必填内容'
        self.assertIn('N',validate_checks(values,checks)[1])
        checks['N']['result']='confirmed'
        self.assertNotIn('N',validate_checks(values,checks)[1])
    @override_settings(DEMO_MODE=False)
    def test_production_import_preview_confirmation(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        self.client.force_login(self.admin)
        r=self.client.post(reverse('import'),{'title':'正式导入流程测试','file':SimpleUploadedFile('students.xlsx',self.workbook())},secure=True)
        self.assertEqual(r.status_code,200);self.assertEqual(ImportPreview.objects.count(),1)
        p=ImportPreview.objects.get();self.assertNotIn('DEMO2026001',p.payload_cipher)
        r=self.client.post(reverse('import'),{'confirm':str(p.pk)},secure=True)
        self.assertEqual(r.status_code,302);self.assertEqual(Batch.objects.count(),2)
        self.assertFalse(ImportPreview.objects.exists())
    def test_public_demo_admin_blocked(self):
        self.assertEqual(self.client.get(reverse('dashboard'),HTTP_X_REAL_IP='203.0.113.10').status_code,403)
    @override_settings(DEMO_MODE=False)
    def test_plain_http_real_import_blocked(self):
        self.client.force_login(self.admin)
        self.assertEqual(self.client.post(reverse('import'),{}).status_code,403)
    def test_other_student_does_not_leak(self):
        self.assertEqual(self.client.get(reverse('reveal',args=['BT'])).json()['value'],'DEMO-GUARDIAN-001')
        self.assertEqual(self.client.get(reverse('student_detail',args=[Student.objects.get(source_row=3,batch=self.batch).pk])).status_code,302)
    def test_template_fullwidth_options_and_original_text_preserved(self):
        rows=demo_rows();rows[0]['values']['I']='港澳居民来往内地通行证（香港）';rows[0]['values']['N']=' 原文（保留全角标点） '
        parsed,_=parse_import(self.workbook(rows))
        self.assertEqual(parsed[0]['values']['I'],rows[0]['values']['I'])
        self.assertEqual(parsed[0]['values']['N'],rows[0]['values']['N'])
        checks=self.all_checks();checks['I']={'result':'confirmed','value':'港澳居民来往内地通行证（香港）'}
        checks['AA']={'result':'confirmed','value':'新住址（保留全角）'}
        s=self.complete(checks)
        self.assertEqual(s.current()['AA'],'新住址（保留全角）')
    def test_import_command_refuses_demo_mode(self):
        with self.assertRaisesMessage(Exception,'Production import requires HTTPS production mode'):
            call_command('import_batch','missing.xlsx',title='测试',activate=True)
