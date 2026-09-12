from collections import Counter
from pathlib import Path

import openpyxl
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from checks.crypto import encrypt, identity
from checks.models import Audit, Batch, Student
from checks.schema import BY_KEY, FIELDS


LABEL_TO_KEY = {field["label"].replace("*", ""): field["key"] for field in FIELDS}
REGION_ALIASES = {
    "北京市海淀区": "北京北京市海淀区",
    "上海市长宁区": "上海上海市长宁区",
    "重庆市长寿区": "重庆重庆市长寿区",
    "重庆市云阳县": "重庆重庆市云阳县",
    "重庆市北碚区": "重庆重庆市北碚区",
}
FIX_FIELDS = {
    "户口所在地行政区划": ("Q", "R"),
    "现住址行政区划": ("Z", "AA"),
    "成员一户口所在地行政区划": ("BO", "BP"),
}


def clean(value):
    if value is None:
        return ""
    return str(value).strip()


class Command(BaseCommand):
    help = "Apply the six address fields from the school's reviewed workbook to students who have not started checking."

    def add_arguments(self, parser):
        parser.add_argument("path")
        parser.add_argument("--apply", action="store_true")
        parser.add_argument("--report")
        parser.add_argument("--allow-started", action="store_true")

    def handle(self, *args, **options):
        path = Path(options["path"]).resolve()
        if not path.is_file() or path.suffix.lower() != ".xlsx":
            raise CommandError("The supplied .xlsx file does not exist.")
        batch = Batch.objects.filter(active=True).first()
        if not batch:
            raise CommandError("There is no active batch.")

        try:
            workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
            sheet = workbook["新生1"]
            fix_sheet = workbook["修复记录"]
        except Exception as exc:
            raise CommandError("无法读取“新生1”或“修复记录”工作表。") from exc

        rows = sheet.iter_rows(values_only=True)
        headers = [clean(value) for value in next(rows, ())]
        label_to_header = {label.replace("*", ""): label for label in headers}
        required = ["姓名*", "身份证件号*"]
        for region_label, (region_key, address_key) in FIX_FIELDS.items():
            required.extend([
                BY_KEY[region_key]["label"],
                BY_KEY[address_key]["label"],
            ])
        missing = [label for label in required if label not in headers]
        if missing:
            raise CommandError("工作表缺少字段：" + "、".join(missing))

        source = []
        for row_number, cells in enumerate(rows, 2):
            if not any(clean(value) for value in cells):
                continue
            row = dict(zip(headers, cells))
            name = clean(row[label_to_header["姓名"]])
            student_id = clean(row[label_to_header["身份证件号"]]).upper()
            source.append((row_number, name, student_id, row))
        workbook.close()

        identities = Counter(identity(name, student_id) for _, name, student_id, _ in source)
        students = {student.identity_hash: student for student in Student.objects.filter(batch=batch)}
        updates = []
        anomalies = []
        counts = Counter()

        for row_number, name, student_id, row in source:
            ident = identity(name, student_id)
            reason = ""
            if not name or not student_id:
                reason = "姓名或身份证号为空"
            elif identities[ident] != 1:
                reason = "姓名与身份证号组合重复"
            elif ident not in students:
                reason = "无法与系统学生精确匹配"
            elif not options["allow_started"] and (students[ident].status != "new" or students[ident].revision != 0):
                reason = "该学生已经开始核对，未自动写入"
            if reason:
                anomalies.append((row_number, name, student_id, "整行", reason, ""))
                counts["整行跳过"] += 1
                continue

            student = students[ident]
            patch = {}
            for field_label, (region_key, address_key) in FIX_FIELDS.items():
                region_label = BY_KEY[region_key]["label"]
                address_label = BY_KEY[address_key]["label"]
                region = clean(row[label_to_header[region_label.replace("*", "")]])
                region = REGION_ALIASES.get(region, region)
                address = clean(row[label_to_header[address_label.replace("*", "")]])
                if region and region not in BY_KEY[region_key]["options"]:
                    anomalies.append((row_number, name, student_id, region_key, "行政区划不在模板选项中", region))
                    counts[f"{region_key}:选项异常"] += 1
                elif region and not address:
                    anomalies.append((row_number, name, student_id, address_key, "行政区划有值但具体地址为空", region))
                    counts[f"{address_key}:地址为空"] += 1
                elif region and len(address) > 500:
                    anomalies.append((row_number, name, student_id, address_key, "具体地址过长", address[:30] + "……"))
                    counts[f"{address_key}:地址过长"] += 1
                else:
                    patch[region_key] = region
                    patch[address_key] = address
            for key, value in patch.items():
                counts[f"{key}:写入"] += 1
            if patch:
                updates.append((student, patch))

        if options["report"]:
            self._write_report(Path(options["report"]), batch, len(source), len(updates), counts, anomalies)
        if options["apply"]:
            with transaction.atomic():
                for student, patch in updates:
                    school = student.school
                    school.update(patch)
                    student.school_cipher = encrypt(school)
                    student.save(update_fields=["school_cipher"])
                Audit.objects.create(
                    actor="system-school-fix",
                    action="应用学校地址户籍修复表",
                    target=str(batch.pk),
                    detail_cipher=encrypt({
                        "source_file": path.name,
                        "source_rows": len(source),
                        "updated_students": len(updates),
                        "field_counts": dict(counts),
                        "anomalies": len(anomalies),
                    }),
                )
        mode = "APPLIED" if options["apply"] else "DRY-RUN"
        self.stdout.write(f"{mode}: source={len(source)} updated={len(updates)} anomalies={len(anomalies)}")
        for key in sorted(counts):
            self.stdout.write(f"{key}={counts[key]}")

    def _write_report(self, path, batch, source_rows, updated_rows, counts, anomalies):
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook = openpyxl.Workbook()
        summary = workbook.active
        summary.title = "导入汇总"
        summary.append(["项目", "数量"])
        for label, value in [
            ("批次", batch.title), ("源表学生数", source_rows), ("将更新学生数", updated_rows),
            ("异常项目数", len(anomalies)),
        ]:
            summary.append([label, value])
        summary.append([])
        summary.append(["字段结果", "数量"])
        for key, value in sorted(counts.items()):
            if ":" in key:
                field, result = key.split(":", 1)
                label = BY_KEY.get(field, {}).get("label", field).replace("*", "")
                summary.append([f"{label}：{result}", value])
            else:
                summary.append([key, value])

        detail = workbook.create_sheet("未写入项目")
        detail.append(["源表行号", "学生姓名", "学生证件号（脱敏）", "字段", "原因", "源值（脱敏）"])
        for row_number, name, student_id, key, reason, value in anomalies:
            masked = student_id[:4] + "**********" + student_id[-4:] if len(student_id) >= 8 else student_id
            detail.append([row_number, name, masked, BY_KEY.get(key, {}).get("label", key).replace("*", ""), reason, value])
        detail.auto_filter.ref = f"A1:F{max(1, detail.max_row)}"
        detail.freeze_panes = "A2"
        for column, width in zip("ABCDEF", [12, 14, 24, 24, 42, 32]):
            detail.column_dimensions[column].width = width
        workbook.save(path)
