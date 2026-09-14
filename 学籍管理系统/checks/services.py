import hashlib
import io
import zipfile
import base64
import re
from datetime import timedelta
import openpyxl
from PIL import Image, UnidentifiedImageError
from openpyxl.drawing.image import Image as ExcelImage
from django.db import transaction
from django.db.models import Max, Prefetch
from django.utils import timezone
from .crypto import encrypt, digest, identity
from .models import Batch, Student, Submission, Audit, Throttle
from .schema import FIELDS, VISIBLE, BY_KEY, SENSITIVE, text_value, validate_checks, validate_value, check_status, check_value

SIGNATURE_RE = re.compile(r"^data:image/png;base64,[A-Za-z0-9+/=]+$")


def validate_signature(signature):
    signature = str(signature or "").strip()
    if not SIGNATURE_RE.fullmatch(signature):
        return "请在签名框中完成手写签字。"
    encoded = signature.split(",", 1)[1]
    if len(signature) > 800000:
        return "签名图片过大，请重新签字。"
    try:
        data = base64.b64decode(encoded, validate=True)
    except (ValueError, TypeError):
        return "签名格式无效，请重新签字。"
    if len(data) < 100 or not data.startswith(b"\x89PNG\r\n\x1a\n"):
        return "签名内容无效，请重新签字。"
    try:
        with Image.open(io.BytesIO(data)) as image:
            if image.format != 'PNG' or not (200 <= image.width <= 2500 and 80 <= image.height <= 1500):
                return "签名图片尺寸无效，请重新签字。"
            rgba = image.convert('RGBA')
            white = Image.new('RGBA', rgba.size, 'white')
            white.alpha_composite(rgba)
            gray = white.convert('L')
            gray.thumbnail((500, 300))
            if sum(pixel < 230 for pixel in gray.getdata()) < 20:
                return "签名笔迹过少，请重新完整签字。"
    except (UnidentifiedImageError, OSError, ValueError):
        return "签名内容无效，请重新签字。"
    return None

def audit(user, action, target, detail=None):
    return Audit.objects.create(actor=str(user), action=action, target=str(target), detail_cipher=encrypt(detail or {}))

@transaction.atomic
def attempt(key, limit=10, seconds=300):
    now = timezone.now()
    bucket, _ = Throttle.objects.get_or_create(key=digest(key), defaults={'started_at': now})
    bucket = Throttle.objects.select_for_update().get(pk=bucket.pk)
    if bucket.started_at < now - timedelta(seconds=seconds):
        bucket.count = 0; bucket.started_at = now
    allowed = bucket.count < limit
    if allowed: bucket.count += 1
    bucket.save()
    return allowed

def parse_import(content):
    if len(content) > 10 * 1024 * 1024: raise ValueError('文件不能超过 10 MB。')
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            if sum(i.file_size for i in z.infolist()) > 80 * 1024 * 1024: raise ValueError('解压后的文件过大。')
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    except Exception as exc: raise ValueError('无法读取 Excel，请上传原模板格式的 .xlsx 文件。') from exc
    if '新生1' not in wb.sheetnames: raise ValueError('缺少工作表“新生1”。')
    ws = wb['新生1']
    rows = ws.iter_rows()
    header = next(rows, ())
    if [c.value for c in header] != [f['label'] for f in FIELDS]: raise ValueError('列名或顺序与原模板不一致，请使用下载的模板。')
    result=[]; errors=[]; serials=set(); ids=set()
    for row_num, cells in enumerate(rows, 2):
        if row_num > 5001: errors.append('单批次最多导入 5000 人。'); break
        if not any(c.value is not None for c in cells): continue
        # Retain source text verbatim, including punctuation and optional-column spacing.
        values={f['key']:(c.value if isinstance(c.value,str) else text_value(c.value)) for f,c in zip(FIELDS,cells)}
        for f,c in zip(FIELDS,cells):
            if c.data_type == 'f': errors.append(f'第 {row_num} 行 {f["label"]}：请粘贴为值，不能导入公式。')
            if f['key'] in {'A','J','Y','BT'} and c.value is not None and c.data_type == 'n':
                errors.append(f'第 {row_num} 行 {f["label"]}：必须使用文本格式，避免号码精度丢失。')
        for k in ['A','B','I','J']:
            if not values[k].strip(): errors.append(f'第 {row_num} 行 {BY_KEY[k]["label"]}：不能为空。')
        for k in ['A','B','V','Y']:
            if len(values[k])>100: errors.append(f'第 {row_num} 行 {BY_KEY[k]["label"]}：内容过长。')
        for k in ['I','J']:
            if values[k]:
                err=validate_value(k,values[k],values)
                if err: errors.append(f'第 {row_num} 行 {BY_KEY[k]["label"]}：{err}')
        serial=digest(values['A']); ident=identity(values['B'],values['J'])
        if serial in serials: errors.append(f'第 {row_num} 行：学籍流水号重复。')
        if ident in ids: errors.append(f'第 {row_num} 行：姓名与证件号码组合重复。')
        serials.add(serial); ids.add(ident)
        result.append({'row':row_num,'values':values})
    wb.close()
    if not result and not errors: errors.append('表格中没有学生数据。')
    if errors: raise ValueError('\n'.join(errors[:80]))
    # Canonical data hash rejects duplicate content even if workbook metadata changes.
    import json
    fingerprint=hashlib.sha256(json.dumps(result,ensure_ascii=False,sort_keys=True).encode()).hexdigest()
    return result,fingerprint

@transaction.atomic
def commit_import(rows, fingerprint, title, user, demo=False):
    batch=Batch.objects.create(title=title, fingerprint=fingerprint, demo=demo)
    students=[]
    for row in rows:
        v=row['values']
        students.append(Student(batch=batch, source_row=row['row'], serial_hash=digest(v['A']), identity_hash=identity(v['B'],v['J']),
                       name=v['B'], class_name=v['V'], original_cipher=encrypt(v), has_issue=not bool(v.get('V'))))
    Student.objects.bulk_create(students)
    audit(user,'导入批次',batch.pk,{'count':len(students),'demo':demo})
    return batch

@transaction.atomic
def save_draft(student_id, revision, checks, fields):
    student=Student.objects.select_for_update().select_related('batch').get(pk=student_id)
    if student.status=='submitted' or not student.batch.active or not student.batch.is_open: raise ValueError('核对已提交或已关闭。')
    if student.revision != revision: raise ValueError('记录已在其他页面更新，请刷新后继续。')
    draft=student.draft; base=student.current()
    for f in fields:
        item=checks.get(f['key'],{})
        old=draft.get(f['key'],{})
        draft[f['key']]={'result':check_status(f,item), 'note':str(item.get('note',''))[:500], 'mode':'direct'}
        if not f['readonly']:
            if 'value' not in item or (f['key'] in SENSITIVE and not item.get('loaded')):
                value=check_value(base,f,old)
            else:
                value=text_value(item.get('value',''))[:500]
            if f['key']=='D': value=value.replace('-','')
            draft[f['key']]['value']=value
    student.draft_cipher=encrypt(draft); student.status='draft'; student.revision+=1
    student.save(update_fields=['draft_cipher','status','revision'])
    return student

@transaction.atomic
def submit(student_id, revision, signature=None):
    student=Student.objects.select_for_update().select_related('batch').get(pk=student_id)
    if student.status=='submitted': return student
    if not student.batch.active or not student.batch.is_open: raise ValueError('本次核对已关闭。')
    if student.revision!=revision: raise ValueError('记录已更新，请刷新并重新确认。')
    checks=student.draft; base=student.current()
    values,errors=validate_checks(base,checks)
    if errors: raise ValueError('还有未完成或格式错误的项目，请返回对应步骤检查。')
    signature_error = validate_signature(signature)
    if signature_error: raise ValueError(signature_error)
    versions=student.submissions.aggregate(n=Max('version'))['n'] or 0
    editable={f['key']:values[f['key']] for f in VISIBLE if not f['readonly']}
    changes={k:{'old':base.get(k,''),'new':v} for k,v in editable.items() if v!=base.get(k,'')}
    Submission.objects.create(student=student,version=versions+1,payload_cipher=encrypt({'values':editable,'checks':checks,'changes':changes,'signature':signature}))
    student.status='submitted'; student.submitted_at=timezone.now(); student.change_count=len(changes)
    student.has_issue=any(
        check_status(field,checks.get(field['key'],{}))=='unconfirmed' or not values.get(field['key'])
        for field in VISIBLE if field['readonly']
    )
    student.revision+=1; student.save()
    return student

@transaction.atomic
def reopen(student_id,user):
    s=Student.objects.select_for_update().get(pk=student_id)
    if s.status != 'submitted': return
    s.status='draft'; s.draft_cipher=encrypt({}); s.revision+=1; s.save()
    audit(user,'重新开放核对',s.pk)

def needs_school_attention(student,payload=None):
    """Return the current class issue without trusting legacy Y-based flags."""
    if not student.has_issue:return False
    school_values={**student.original,**student.school}
    if not text_value(school_values.get('V','')):return True
    if payload is None:
        latest=student.submissions.order_by('-version').first()
        payload=latest.payload if latest else {}
    if not payload:return False
    class_check=payload.get('checks',{}).get('V')
    return bool(class_check) and check_status(BY_KEY['V'],class_check)=='unconfirmed'


def public_review_issues(student):
    """Return non-sensitive issue labels suitable for the public progress rosters."""
    submissions = getattr(student, 'progress_submissions', None)
    if submissions is None:
        submissions = list(student.submissions.order_by('version'))
    latest_payload = submissions[-1].payload if submissions else {}
    issues = []
    if needs_school_attention(student, latest_payload):
        issues.append('班级信息待核实')

    legacy_issue_version = None
    submitted_y_versions = []
    for submission in submissions:
        payload = submission.payload
        values = payload.get('values', {})
        y_check = payload.get('checks', {}).get('Y', {})
        if 'Y' in values:
            submitted_y_versions.append(submission.version)
        elif y_check.get('result') in {'incorrect', 'unconfirmed'}:
            legacy_issue_version = submission.version
    if legacy_issue_version is not None and not any(
        version > legacy_issue_version for version in submitted_y_versions
    ):
        current_y = student.draft.get('Y', {})
        if check_status(BY_KEY['Y'], current_y) != 'confirmed':
            issues.append('全国学籍号待重新核对')
    return issues

def safe_cell(cell,value):
    cell.value=value
    if isinstance(value,str): cell.data_type='s'; cell.number_format='@'

def export_workbook(batch,kind,class_name=None):
    students=Student.objects.filter(batch=batch)
    if class_name is not None: students=students.filter(class_name=class_name)
    if kind=='pending': students=students.exclude(status='submitted')
    if kind in {'changes','confirmations'}:
        students=(students.filter(submissions__isnull=False).distinct()
                  .prefetch_related(Prefetch('submissions',queryset=Submission.objects.order_by('version'),to_attr='export_submissions')))
    else:
        students=students.prefetch_related('submissions')
    if kind in {'final','review','pending'}:
        wb=openpyxl.load_workbook(__import__('django').conf.settings.BASE_DIR/'assets/template.xlsx')
        ws=wb['新生1']
        for row,s in enumerate(students,2):
            # Draft values are never exported, but a reopened student's last submitted values remain final.
            values=s.current()
            for col,f in enumerate(FIELDS,1): safe_cell(ws.cell(row,col),values.get(f['key'],''))
        ws.freeze_panes='C2'; ws.auto_filter.ref=f'A1:CH{max(1,len(students)+1)}'
    elif kind=='changes':
        wb=openpyxl.Workbook();ws=wb.active;ws.title='修改明细'
        ws.append(['班级','姓名','提交版本','提交时间','字段','原始值','提交值','核对结果','说明'])
        for s in students:
            original=s.original;school=s.school
            for sub in s.export_submissions:
                p=sub.payload
                for f in VISIBLE:
                    k=f['key']; check=p['checks'].get(k,{})
                    historical_readonly=f['readonly'] or (k=='Y' and k not in p.get('values',{}))
                    if k not in p['changes'] and not (historical_readonly and check_status(f,check)=='unconfirmed'): continue
                    change=p['changes'].get(k,{'old':original.get(k,''),'new':school.get(k,original.get(k,''))})
                    row=[s.class_name,s.name,str(sub.version),timezone.localtime(sub.created_at).strftime('%Y-%m-%d %H:%M:%S'),f['label'].replace('*',''),change['old'],change['new'],'学校待处理' if historical_readonly else '已更正',check.get('note','')]
                    n=ws.max_row+1
                    for col,v in enumerate(row,1): safe_cell(ws.cell(n,col),v)
        ws.freeze_panes='A2'
    else:
        wb=openpyxl.Workbook();ws=wb.active;ws.title='逐项确认'
        ws.append(['班级','姓名','提交版本','提交时间','字段','是否必填','维护方','确认状态','学校原值','家长提交值','是否修改','问题说明'])
        signatures=wb.create_sheet('家长签名')
        signatures.append(['班级','姓名','提交版本','提交时间','手写签名'])
        image_streams=[]
        confirmation_row=2
        signature_row=2
        for s in students:
            original={**s.original,**s.school}
            for sub in s.export_submissions:
                payload=sub.payload;checks=payload.get('checks',{});submitted=payload.get('values',{})
                submitted_at=timezone.localtime(sub.created_at).strftime('%Y-%m-%d %H:%M:%S')
                for f in VISIBLE:
                    key=f['key'];check=checks.get(key,{})
                    historical_readonly=f['readonly'] or (key=='Y' and key not in submitted)
                    old=text_value(original.get(key,''))
                    new=old if historical_readonly else text_value(submitted.get(key,old))
                    row=[s.class_name,s.name,str(sub.version),submitted_at,f['label'].replace('*',''),
                         '是' if f['required'] else '否','学校' if historical_readonly else '家长',
                         '已确认' if check_status(f,check)=='confirmed' else '未确认',old,new,
                         '是' if not historical_readonly and new!=old else '否',str(check.get('note',''))]
                    for col,value in enumerate(row,1): safe_cell(ws.cell(confirmation_row,col),value)
                    confirmation_row+=1
                signature=payload.get('signature','')
                signatures.append([s.class_name,s.name,str(sub.version),submitted_at,''])
                if signature and ',' in signature:
                    try:
                        stream=io.BytesIO(base64.b64decode(signature.split(',',1)[1],validate=True));image_streams.append(stream)
                        picture=ExcelImage(stream);picture.width=220;picture.height=88
                        signatures.add_image(picture,f'E{signature_row}');signatures.row_dimensions[signature_row].height=70
                    except (ValueError,TypeError,OSError):
                        signatures.cell(signature_row,5,'签名图片无法读取')
                else:
                    signatures.cell(signature_row,5,'未保存签名')
                signature_row+=1
        ws.freeze_panes='A2';signatures.freeze_panes='A2'
        widths=[12,12,10,20,24,10,10,10,28,28,10,32]
        for index,width in enumerate(widths,1):ws.column_dimensions[openpyxl.utils.get_column_letter(index)].width=width
        for column,width in {'A':12,'B':12,'C':10,'D':20,'E':34}.items():signatures.column_dimensions[column].width=width
    stream=io.BytesIO();wb.save(stream);return stream.getvalue()
