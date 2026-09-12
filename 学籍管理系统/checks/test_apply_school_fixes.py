import tempfile
from pathlib import Path

import openpyxl
from django.core.management import call_command
from django.test import TestCase

from .crypto import digest, encrypt, identity
from .models import Batch, Student


class ApplySchoolFixesTests(TestCase):
    def test_applies_only_reviewed_address_fields(self):
        student_id = "430528201311260091"
        batch = Batch.objects.create(title="修复测试", fingerprint="f" * 64, active=True)
        original = {
            "A": "serial", "B": "测试学生", "I": "居民身份证", "J": student_id,
            "Q": "旧区划", "R": "旧地址", "Z": "旧现住区划", "AA": "旧现住地址",
            "BO": "旧成员区划", "BP": "旧成员地址",
        }
        student = Student.objects.create(
            batch=batch, source_row=2, serial_hash=digest("serial"),
            identity_hash=identity("测试学生", student_id), name="测试学生",
            original_cipher=encrypt(original),
        )
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "新生1"
        sheet.append([field + "*" for field in ["姓名", "身份证件号", "户口所在地行政区划", "户口所在地具体地址", "现住址行政区划", "现住址具体地址", "成员一户口所在地行政区划", "成员一户口所在地具体地址"]])
        sheet.append(["测试学生", student_id, "湖南省长沙市岳麓区", "湖南省长沙市岳麓区修复地址", "湖南省长沙市雨花区", "湖南省长沙市雨花区现住址", "湖南省长沙市天心区", "湖南省长沙市天心区成员地址"])
        fixes = workbook.create_sheet("修复记录")
        fixes.append(["原表行号", "班级", "姓名", "字段", "原值", "修复后", "修复依据", "置信度"])
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "fix.xlsx"
            workbook.save(path)
            call_command("apply_school_fixes", str(path), apply=True)
        student.refresh_from_db()
        self.assertEqual(student.school["Q"], "湖南省长沙市岳麓区")
        self.assertEqual(student.school["R"], "湖南省长沙市岳麓区修复地址")
        self.assertEqual(student.school["AA"], "湖南省长沙市雨花区现住址")
        self.assertEqual(student.school["BP"], "湖南省长沙市天心区成员地址")
        self.assertEqual(student.original["R"], "旧地址")

    def test_normalizes_direct_administration_regions_and_can_allow_started(self):
        student_id = "430528201311260091"
        batch = Batch.objects.create(title="修复测试", fingerprint="a" * 64, active=True)
        student = Student.objects.create(
            batch=batch, source_row=2, serial_hash=digest("serial"),
            identity_hash=identity("测试学生", student_id), name="测试学生",
            original_cipher=encrypt({"B": "测试学生", "I": "居民身份证", "J": student_id}),
            status="draft", revision=2,
        )
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "新生1"
        labels = ["姓名", "身份证件号", "户口所在地行政区划", "户口所在地具体地址", "现住址行政区划", "现住址具体地址", "成员一户口所在地行政区划", "成员一户口所在地具体地址"]
        sheet.append([label + "*" for label in labels])
        sheet.append(["测试学生", student_id, "北京市海淀区", "北京市海淀区地址", "", "", "", ""])
        workbook.create_sheet("修复记录")
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "fix.xlsx"
            workbook.save(path)
            call_command("apply_school_fixes", str(path), apply=True, allow_started=True)
        student.refresh_from_db()
        self.assertEqual(student.school["Q"], "北京北京市海淀区")
        self.assertEqual(student.status, "draft")
