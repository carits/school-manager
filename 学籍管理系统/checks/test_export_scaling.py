import base64
import io
from unittest.mock import patch

import openpyxl
from PIL import Image, ImageDraw
from openpyxl.worksheet.worksheet import Worksheet
from django.db import connection
from django.test import TestCase
from django.test.utils import CaptureQueriesContext

from .crypto import encrypt
from .models import Batch, Student, Submission
from .schema import FIELDS, VISIBLE
from .services import export_workbook


def signature_data_url():
    image = Image.new('RGB', (300, 120), 'white')
    draw = ImageDraw.Draw(image)
    draw.line([(25, 85), (80, 30), (130, 90), (185, 25), (265, 80)], fill='#142f47', width=5)
    stream = io.BytesIO()
    image.save(stream, format='PNG')
    return 'data:image/png;base64,' + base64.b64encode(stream.getvalue()).decode()


class ExportScalingTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.batch = Batch.objects.create(title='导出扩展测试', fingerprint='export-scaling', active=True)
        cls.signature = signature_data_url()
        cls.students = []
        for index in range(1, 7):
            values = {field['key']: '' for field in FIELDS}
            values.update({'A': f'SERIAL-{index}', 'B': f'学生{index}', 'V': '2601', 'Y': f'G{index:018d}'})
            cls.students.append(Student.objects.create(
                batch=cls.batch,
                source_row=index + 1,
                serial_hash=f'serial-{index}',
                identity_hash=f'identity-{index}',
                name=values['B'],
                class_name=values['V'],
                original_cipher=encrypt(values),
                school_cipher=encrypt({}),
            ))

        # Insert versions out of order so the export's ordered prefetch is exercised.
        cls._submission(cls.students[0], 2, cls.signature)
        cls._submission(cls.students[0], 1, cls.signature, legacy=True)
        cls._submission(cls.students[1], 1, cls.signature)
        cls._submission(cls.students[2], 1, 'data:image/png;base64,' + base64.b64encode(b'not a png').decode())
        cls._submission(cls.students[3], 1, cls.signature)
        # students[4] and students[5] deliberately have no submissions.

    @classmethod
    def _submission(cls, student, version, signature, legacy=False):
        original_name = student.name
        new_name = f'{student.name}-第{version}版'
        checks = {
            field['key']: {'result': 'correct' if legacy else 'confirmed'}
            for field in VISIBLE
        }
        payload = {
            'values': {'B': new_name},
            'checks': checks,
            'changes': {'B': {'old': original_name, 'new': new_name}},
            'signature': signature,
        }
        Submission.objects.create(
            student=student,
            version=version,
            payload_cipher=encrypt(payload),
        )

    def export_with_query_count(self, kind):
        with CaptureQueriesContext(connection) as queries:
            content = export_workbook(self.batch, kind)
        return content, len(queries)

    def test_confirmations_export_multiple_students_versions_and_signatures(self):
        original_max_row = Worksheet.max_row.fget
        max_row_reads = 0

        def tracked_max_row(sheet):
            nonlocal max_row_reads
            max_row_reads += 1
            return original_max_row(sheet)

        with patch.object(Worksheet, 'max_row', property(tracked_max_row)):
            content, query_count = self.export_with_query_count('confirmations')
        self.assertLessEqual(query_count, 2)
        self.assertLess(max_row_reads, 20)

        workbook = openpyxl.load_workbook(io.BytesIO(content))
        confirmations = workbook['逐项确认']
        signatures = workbook['家长签名']
        submission_count = 5
        self.assertEqual(confirmations.max_row, 1 + submission_count * len(VISIBLE))
        self.assertEqual(signatures.max_row, 1 + submission_count)
        self.assertEqual(len(signatures._images), 4)
        self.assertEqual(signatures['C2'].value, '1')
        self.assertEqual(signatures['C3'].value, '2')
        self.assertEqual(signatures['E5'].value, '签名图片无法读取')
        self.assertEqual(confirmations['H2'].value, '已确认')

    def test_changes_export_uses_constant_query_count_and_skips_empty_students(self):
        content, query_count = self.export_with_query_count('changes')
        self.assertLessEqual(query_count, 2)

        workbook = openpyxl.load_workbook(io.BytesIO(content))
        changes = workbook['修改明细']
        self.assertEqual(changes.max_row, 6)
        self.assertEqual([changes.cell(row, 3).value for row in range(2, 4)], ['1', '2'])
        exported_names = {changes.cell(row, 2).value for row in range(2, changes.max_row + 1)}
        self.assertNotIn(self.students[4].name, exported_names)
        self.assertNotIn(self.students[5].name, exported_names)
